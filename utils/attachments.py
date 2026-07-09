"""Validation + text extraction for student-uploaded non-image files.

Pure functions over ``(filename, bytes)`` — no Flask, no DB. Sibling to
``utils.uploads`` (images). Every supported type is extracted to plain text so
the tutor consumes it uniformly across providers and the per-message character
budget can be enforced (the real cost control — see the design spec).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

# ext -> kind. TSV shares the CSV reader with a tab delimiter.
ALLOWED_FILE_EXTS: dict[str, str] = {
    ".csv": "csv",
    ".tsv": "tsv",
    ".xlsx": "xlsx",
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
}

MAX_FILE_BYTES = 5 * 1024 * 1024      # 5 MB coarse gate (real cost control is the char budget)
MAX_EXTRACTED_CHARS = 15000           # ~4k tokens across all files in one message


class AttachmentValidationError(ValueError):
    """Bad type or oversize file."""


class AttachmentExtractionError(ValueError):
    """File matched a known type but could not be parsed (corrupt / missing parser)."""


class EmptyExtractionError(ValueError):
    """File parsed but yielded no usable text (e.g. scanned image-only PDF)."""


@dataclass(frozen=True)
class ValidatedAttachment:
    """An accepted upload: original filename, kind, extracted text, raw bytes."""

    filename: str
    kind: str
    extracted_text: str
    data: bytes

    @property
    def size_bytes(self) -> int:
        return len(self.data)


def _kind_for(filename: str) -> str | None:
    return ALLOWED_FILE_EXTS.get(Path(filename or "").suffix.lower())


def _extract_delimited(data: bytes, delimiter: str) -> str:
    text = data.decode("utf-8", errors="replace")
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        return "\n".join(", ".join(row) for row in reader)
    except Exception as exc:
        raise AttachmentExtractionError(f"Could not read delimited file: {exc}") from exc


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _extract_xlsx(data: bytes) -> str:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise AttachmentExtractionError("openpyxl is required to read .xlsx files.") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise AttachmentExtractionError(f"Could not read spreadsheet: {exc}") from exc
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [("" if c is None else str(c)) for c in row]
            if any(cells):
                lines.append(", ".join(cells))
    return "\n".join(lines)


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise AttachmentExtractionError("python-docx is required to read .docx files.") from exc
    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise AttachmentExtractionError(f"Could not read document: {exc}") from exc
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_pdf(data: bytes) -> str:
    try:
        import pypdf
    except ImportError as exc:  # pragma: no cover
        raise AttachmentExtractionError("pypdf is required to read .pdf files.") from exc
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise AttachmentExtractionError(f"Could not read PDF: {exc}") from exc


_EXTRACTORS = {
    "csv": lambda d: _extract_delimited(d, ","),
    "tsv": lambda d: _extract_delimited(d, "\t"),
    "txt": _extract_txt,
    "xlsx": _extract_xlsx,
    "docx": _extract_docx,
    "pdf": _extract_pdf,
}


def validate_file(filename: str, data: bytes) -> ValidatedAttachment:
    """Validate + extract one upload. Raises on bad type, oversize, or empty text."""
    if not data:
        raise AttachmentValidationError(f"'{filename}' is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentValidationError(
            f"'{filename}' is {len(data)} bytes; max is {MAX_FILE_BYTES} bytes."
        )
    kind = _kind_for(filename)
    if kind is None:
        raise AttachmentValidationError(
            f"'{filename}' is not a supported file type "
            f"({', '.join(sorted(ALLOWED_FILE_EXTS))})."
        )
    text = _EXTRACTORS[kind](data).strip()
    if not text:
        raise EmptyExtractionError(
            f"'{filename}' contained no readable text (a scanned/image-only file "
            "can't be read — attach the text or a screenshot instead)."
        )
    return ValidatedAttachment(filename=filename, kind=kind, extracted_text=text, data=data)


def validate_files(items: list[tuple[str, bytes]]) -> list[ValidatedAttachment]:
    """Validate each file, then truncate combined text to MAX_EXTRACTED_CHARS."""
    atts = [validate_file(name, data) for (name, data) in items]
    budget = MAX_EXTRACTED_CHARS
    out: list[ValidatedAttachment] = []
    for att in atts:
        if budget <= 0:
            text = f"\n[…truncated {len(att.extracted_text)} chars for length…]"
        elif len(att.extracted_text) > budget:
            dropped = len(att.extracted_text) - budget
            text = att.extracted_text[:budget] + f"\n[…truncated {dropped} chars for length…]"
            budget = 0
        else:
            text = att.extracted_text
            budget -= len(att.extracted_text)
        out.append(
            ValidatedAttachment(att.filename, att.kind, text, att.data)
            if text != att.extracted_text
            else att
        )
    return out


def attachments_to_text_block(atts: list["ValidatedAttachment"]) -> str:
    """Render attachments as labeled text blocks appended to a student message."""
    if not atts:
        return ""
    return "".join(f"\n\n[Attachment: {a.filename}]\n{a.extracted_text}" for a in atts)
