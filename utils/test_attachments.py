import io
import pytest
from utils import attachments as A


def _xlsx_bytes(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(paragraphs):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_csv_extracts_cells():
    att = A.validate_file("t.csv", b"a,b\n1,2\n")
    assert att.kind == "csv"
    assert "a" in att.extracted_text and "2" in att.extracted_text


def test_tsv_extracts_cells():
    att = A.validate_file("t.tsv", b"a\tb\n1\t2\n")
    assert att.kind == "tsv"
    assert "1" in att.extracted_text


def test_txt_decodes():
    att = A.validate_file("n.txt", "hello world".encode("utf-8"))
    assert att.extracted_text.strip() == "hello world"


def test_xlsx_extracts_cells():
    att = A.validate_file("s.xlsx", _xlsx_bytes([["h1", "h2"], [10, 20]]))
    assert att.kind == "xlsx"
    assert "h1" in att.extracted_text and "20" in att.extracted_text


def test_docx_extracts_paragraphs():
    att = A.validate_file("d.docx", _docx_bytes(["first para", "second para"]))
    assert "first para" in att.extracted_text and "second para" in att.extracted_text


def test_unknown_type_rejected():
    with pytest.raises(A.AttachmentValidationError):
        A.validate_file("x.exe", b"MZ\x90\x00")


def test_oversize_rejected():
    with pytest.raises(A.AttachmentValidationError):
        A.validate_file("t.txt", b"x" * (A.MAX_FILE_BYTES + 1))


def test_empty_extraction_rejected():
    # a .txt that decodes to only whitespace has no usable content
    with pytest.raises(A.EmptyExtractionError):
        A.validate_file("blank.txt", b"   \n  ")


def test_budget_truncates_combined_text():
    big = ("row,val\n" + "x,1\n" * 20000).encode("utf-8")  # well over 15000 chars
    atts = A.validate_files([("big.csv", big)])
    total = sum(len(a.extracted_text) for a in atts)
    assert total <= A.MAX_EXTRACTED_CHARS + 60  # + room for the marker
    assert "truncated" in atts[-1].extracted_text


def test_text_block_format():
    atts = A.validate_files([("t.csv", b"a,b\n1,2\n")])
    block = A.attachments_to_text_block(atts)
    assert block.startswith("\n\n[Attachment: t.csv]\n")


def test_corrupt_csv_raises_extraction_error():
    bad = bytes(range(256)) * 100  # binary junk with a .csv name
    with pytest.raises(A.AttachmentExtractionError):
        A.validate_file("corrupt.csv", bad)
